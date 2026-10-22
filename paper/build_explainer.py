"""Build paper/EDIA_explained.docx -- the paper explained section by section, in plain
language AND technical detail, so that every teammate can understand and defend it.

    python paper/build_explainer.py

Every number here is the paper's own (paper/edia_ieee.tex; sources in paper/NUMBERS.md).
If a result in the paper changes, update it here too.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from build_docx import TBLPR_ORDER, TCPR_ORDER, insert_ordered  # noqa: E402  (shared XML helpers)

OUT = HERE / "EDIA_explained.docx"
FIG = HERE / "fig_pipeline.png"
BODY_FONT = "Calibri"
SIMPLE_FILL = "EAF2FB"       # light blue: "In simple terms" boxes
HEAD_FILL = "D9E2F3"         # table header rows
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
AUTHORS = "Jaitra Shah, Manan Mehta, Soham Dwivedi, Aayusman Pati"
AFFIL = ("Dept. of Computer Science and Business Systems, Mukesh Patel School of Technology "
         "Management and Engineering, NMIMS, Mumbai, India")
PAPER_TITLE = ("Where Meeting Reasoning Gets Lost: A Feasibility Study of Decision-Context Gap "
               "Detection in Enterprise Meetings")

doc = Document()


# ----------------------------------------------------------------------------- helpers
def setup():
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)        # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = BODY_FONT, Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True
    fp = sec.footer.paragraphs[0]                                # footer: "Page N"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(9)
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        r = fp.add_run()
        r.font.size = Pt(9)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)


def rich(p, text: str):
    """**bold** and *italic* markup inside a string."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)
    return p


def para(text: str = "", style=None, after=None, align=None):
    p = doc.add_paragraph(style=style)
    rich(p, text)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        rich(p, it)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(cell._tc.get_or_add_tcPr(), shd, TCPR_ORDER)


def simple_box(*paragraphs: str, title: str = "In simple terms"):
    """Shaded call-out (one-cell table). Lines starting with '- ' become bullets."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.6)
    shade(cell, SIMPLE_FILL)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge == "left":
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "24"); el.set(qn("w:color"), "2E75B6")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    insert_ordered(t._tbl.tblPr, borders, TBLPR_ORDER)
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 90), ("left", 170), ("bottom", 90), ("right", 170)):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa"); mar.append(el)
    insert_ordered(t._tbl.tblPr, mar, TBLPR_ORDER)
    head = cell.paragraphs[0]
    r = head.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    head.paragraph_format.space_after = Pt(3)
    head.paragraph_format.keep_with_next = True
    for txt in paragraphs:
        if txt.startswith("- "):
            p = cell.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            rich(p, txt[2:])
        else:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            rich(p, txt)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def technical(label: str = "Technical detail"):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = ACCENT
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True


def table(header, rows, widths_cm, font=9.5, align=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    align = align or ["l"] * len(header)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT, "c": WD_ALIGN_PARAGRAPH.CENTER}
    for ri, row in enumerate([header] + rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri].cells[ci]
            cell.width = Cm(widths_cm[ci])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = ri < len(rows)
            p.alignment = amap[align[ci]] if ri else WD_ALIGN_PARAGRAPH.LEFT
            rich(p, str(val))
            for run in p.runs:
                run.font.size = Pt(font)
                if ri == 0:
                    run.bold = True
            if ri == 0:
                shade(cell, HEAD_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(text):
    p = para(text, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.size = Pt(9)
        r.italic = True


def qa(question: str, answer: str):
    p = para(f"**Q. {question}**", after=1)
    p.paragraph_format.keep_with_next = True
    para(answer, after=7)


# ----------------------------------------------------------------------------- content
setup()

title = doc.add_paragraph(style="Title")
title.add_run("EDIA, explained")
sub = doc.add_paragraph(style="Subtitle")
sub.add_run("Understanding our paper, section by section, in simple terms and in technical detail")
para(f"**Paper:** {PAPER_TITLE}", after=2)
para(f"**Authors:** {AUTHORS} · {AFFIL}", after=2)
para("**EDIA** stands for **Explainable Decision Intelligence Assistant**: a system that reads a "
     "meeting and explains what was decided and why.", after=10)

# 0 -------------------------------------------------------------------------------------------
doc.add_heading("How to read this document", level=1)
para("Every section has two layers. The shaded **In simple terms** box explains the idea with no "
     "background needed. The **Technical detail** that follows is what you need to answer an examiner or "
     "a reviewer. Read only the boxes for a 15-minute overview; read everything to be able to defend "
     "the paper. Section numbers here follow the paper's own order, and Section 12 maps every table "
     "in the paper to the place it is explained.")

doc.add_heading("The paper in one paragraph", level=2)
simple_box(
    "Meetings end with decisions, but the minutes usually record only **what** was decided, not **why**. "
    "We asked whether today's AI can (1) dig the “why” out of a meeting transcript and (2) notice "
    "when the minutes left it out. We built a system to do both, called **EDIA**, and tested it carefully.",
    "- **Spotting what the minutes left out:** a simple word-overlap check beats two smarter-sounding AI "
    "methods. The “smart” ones keep missing things that *are* in the minutes.",
    "- **Digging out the reasons:** EDIA finds roughly **a third to over half** of the reasons a person "
    "finds (33–58%). Most of what it misses is reasoning that is spread far apart in the meeting or "
    "never said outright.",
    "- **Other well-known methods do not do better.** Asking the AI to think step by step, Microsoft's "
    "GraphRAG, and multi-step search all stay at or below EDIA.",
    "So the honest answer to “is this feasible?” is **partially**, and the paper explains exactly "
    "where and why it stops working.")

# 1 -------------------------------------------------------------------------------------------
doc.add_heading("1. The problem: where meeting reasoning gets lost", level=1)
simple_box(
    "A real example from our test data: a product team designing a TV remote discusses how to power it. "
    "They consider a battery, solar power and a dynamo, and pick an unusual option, kryptonite, reasoning "
    "that the chip is the most expensive part anyway. Nobody ever finds out what kryptonite costs, and "
    "the Industrial Designer would rather have used a battery. The summary keeps one line: "
    "*“Kryptonite was chosen as the preferred energy source.”*",
    "Later, someone asks: why not a battery? Who disagreed? Was the price ever checked? All of that was "
    "said in the meeting, and none of it is in the record. Each of those missing pieces is what we call "
    "a **decision-context gap**. The project asks whether a computer can find them automatically.")
technical()
para("A **decision-context gap** is a reasoning element voiced in a meeting but absent from the filed "
     "document. Reasoning elements are typed: the **decision** itself, **alternatives** considered, "
     "**rationale**, **conditions** attached, **dissent**, and the **owner/approver** of the decision. "
     "Detecting gaps needs two capabilities: **extraction** (find the reasoning a meeting actually "
     "contains) and **gap detection** (decide whether each element is present in the document). Both look "
     "within reach of modern language models. The paper tests that assumption instead of taking it for "
     "granted, which is why it is framed as a **feasibility study**.")

# 2 -------------------------------------------------------------------------------------------
doc.add_heading("2. What the paper contributes", level=1)
simple_box(
    "The paper gives four things: a working system (EDIA), a test set to measure such systems fairly "
    "(edia_db), an explanation of *why* the system misses what it misses, and a fair comparison against "
    "three popular alternatives.")
technical()
table(["Contribution", "What it is", "Where in the paper"],
      [["EDIA", "End-to-end pipeline for structured decision-reasoning extraction + gap detection, "
        "run entirely on open models (no paid APIs)", "Sec. III, Fig. 1"],
       ["edia_db benchmark", "371 transcript–document pairs, 6,647 labelled elements, in three tiers "
        "with explicit label provenance", "Sec. IV, Table I"],
       ["Failure diagnosis", "Every residual miss read by hand and categorised; multi-hop reasoning "
        "dominates", "Sec. VI, Table V"],
       ["Comparative analysis", "Chain-of-thought, GraphRAG and multi-hop retrieval run on the same "
        "meetings and scored the same way", "Sec. VII, Tables VI–VII"]],
      [3.4, 9.6, 3.6], font=9)

# 3 -------------------------------------------------------------------------------------------
doc.add_heading("3. How EDIA works (paper Sec. III)", level=1)
simple_box(
    "EDIA works like a careful note-taker in two stages.",
    "- **Stage 1, extraction.** The meeting transcript is cut into bite-sized pieces. For each piece, an "
    "AI model fills in a fixed form: what was decided, what else was considered, why, under what "
    "conditions, who disagreed, and the exact quotes that prove it.",
    "- **Stage 2, gap detection.** Each reason found is checked against the minutes: is it in there or not?",
    "The AI is a free, open model (**Mistral 7B**) that runs on an ordinary computer, so no meeting data "
    "ever leaves the machine.")
if FIG.exists():
    doc.add_picture(str(FIG), width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption("Fig. 1 of the paper: the EDIA pipeline. Grey = extraction; dashed = evaluation.")
technical("Technical detail: extraction")
bullets([
    "**Schema.** Each decision becomes a record: decision statement, topic, status (decided, provisional, "
    "deferred, rejected, unclear), alternatives (with a rejected flag and reason), rationale, conditions, "
    "dissent, owner, approver, and verbatim evidence quotes with speaker labels.",
    "**Chunking.** Transcripts are split into chunks of at most 1,200 tokens with a two-turn overlap, so a "
    "sentence cut at a boundary still appears whole in one chunk.",
    "**Prompting.** The prompt forbids inferring anything not in the excerpt and requires quotes to be "
    "copied character-for-character.",
    "**Deterministic settings.** Temperature 0, top-k 1, fixed seed, 8,192-token context, 2,048-token "
    "output cap, JSON-constrained output; one automatic repair call if the output does not parse. "
    "“Deterministic” means the model always takes its single most likely next word, so runs are "
    "(almost) repeatable.",
])
technical("Technical detail: gap detection")
para("A detector gives each (document, element) pair a **present-score** between 0 and 1; below the "
     "threshold of **0.55** the element is called a gap. Four detectors are compared:")
table(["Detector", "In simple terms", "Technically"],
      [["Lexical", "Do the same words appear in the minutes?", "Token-level F1 overlap between element and document"],
       ["NLI", "Does the minutes text *logically imply* this reason?", "Entailment probability from roberta-large-mnli, "
        "a model trained to judge whether one text follows from another"],
       ["LLM judge", "Ask the AI: “is this point in the minutes?”", "Mistral 7B returns a verdict + confidence"],
       ["Ensemble", "Call it present if any of the three says present", "Each member's score is shifted by its own "
        "threshold (so 0.5 = “present” for all), then the maximum is taken"]],
      [2.4, 6.0, 8.2], font=9)

# 4 -------------------------------------------------------------------------------------------
doc.add_heading("4. The test data: edia_db (paper Sec. IV, Table I)", level=1)
simple_box(
    "To grade the system we need meetings where the right answers are known. We built three kinds:",
    "- **Gold:** 5 real meetings where a person marked every reason by hand. Most trustworthy, but small.",
    "- **Corpus:** 227 more real meetings, labelled automatically by an AI model. Large, but the labels "
    "are only “silver”: good guesses, not checked answers.",
    "- **Synthetic:** 139 made-up meetings where we decided in advance which reasons the minutes would "
    "keep and which they would drop, so the answers are known by design.",
    "Each tier covers a weakness of the others: gold is reliable but tiny, corpus is real and large but "
    "auto-labelled, synthetic has known answers but is cleaner than real speech.")
technical()
table(["Tier", "Pairs", "Elements", "Gap rate", "Labels come from"],
      [["Gold (QMSum)", "5", "50", "0.76", "Human annotation against transcript spans"],
       ["Corpus (QMSum)", "227", "4,930", "0.93", "The NLI detector at threshold 0.55 (silver)"],
       ["Synthetic", "139", "1,667", "0.59", "Constructed: each element's fate fixed before any text is written"],
       ["**Total**", "**371**", "**6,647**", "0.84", ""]],
      [3.0, 1.5, 1.8, 1.8, 8.5], font=9, align=["l", "r", "r", "r", "l"])
bullets([
    "**Source of real meetings.** QMSum, a public benchmark of meeting transcripts with human-written "
    "summaries from three domains: Academic (research-group meetings, ICSI corpus), Product (design "
    "meetings, AMI corpus) and Committee (parliamentary committees). We treat QMSum's summaries as the "
    "“filed document”. The corpus tier covers Academic 59, Committee 36 and Product 132 meetings.",
    "**Gap rate** = share of elements that are missing from the document. 0.93 in the corpus means 93% "
    "of elements are gaps, which matters a lot for how scores behave (Section 5).",
    "**Why scoring NLI on the corpus tier is not allowed.** The corpus labels *were produced by* NLI, so "
    "NLI would just be agreeing with itself. The paper excludes that comparison as circular.",
    "**Synthetic generation.** Mistral 7B, 79 topics covering 8 organisational domains × 10 decision "
    "archetypes (e.g. vendor selection, build-vs-buy, scope cut), with an automated check that "
    "“omitted” content really did not leak into the document. It is the only tier whose gap rate "
    "is near 50/50, which is why detectors are compared there.",
])

doc.add_heading("The human spot-check of the silver labels", level=2)
simple_box(
    "Since the 227 corpus meetings were labelled by an AI, one of us checked a sample by hand, without "
    "looking at the AI's answers first (“blind”). Result: whenever a point really was missing, the "
    "AI labels agreed. But about **half** of what they called “missing” was actually in the summary, "
    "just reworded. So the corpus overstates how much gets lost.")
technical()
para("One annotator judged **52 elements** from three randomly sampled corpus pairs (two Academic, one "
     "Committee), blind to the silver labels. Gap recall 1.0 (no real gap was missed), but **24 of the 46 "
     "predicted gaps** were judged present, 23 of them as paraphrases: gap precision **0.478**, present "
     "recall **0.20**, balanced accuracy **0.60**. The direction of the error is the same in all three "
     "pairs, although 19 of the 24 disagreements come from the Committee pair. Consequence: corpus-tier "
     "gap rates should be read as **upper bounds**.")

# 5 -------------------------------------------------------------------------------------------
doc.add_heading("5. How the results are measured, and why some numbers mislead", level=1)
para("This section is the key to reading every table in the paper.")

doc.add_heading("5.1 The base-rate problem, and balanced accuracy", level=2)
simple_box(
    "Suppose 93 out of 100 points are missing from the minutes. A lazy detector that always answers "
    "“missing” is right 93% of the time and gets an F1 score of about 0.96, while doing nothing "
    "useful at all. A big number can therefore hide a useless detector.",
    "**Balanced accuracy** fixes this: it averages how well the detector does on the “missing” "
    "points and on the “present” points separately. The lazy detector scores exactly **0.5**, "
    "i.e. a coin flip. That is why the paper leads with balanced accuracy, not plain accuracy or F1.")
technical()
bullets([
    "**Balanced accuracy** = (gap recall + present recall) / 2; 0.5 for any constant detector.",
    "**Gap F1** and **present F1** are reported side by side: present F1 shows whether the detector can "
    "recognise content that survived, which is exactly where detectors fail.",
    "**Real example from the paper.** On Committee meetings the lexical detector has the *highest* gap F1 "
    "of any group (0.988) but the *lowest* balanced accuracy (0.532): their summaries are almost entirely "
    "meta-description, so guessing “gap” wins.",
    "This is why detector comparisons are made on the near-balanced synthetic tier.",
])

doc.add_heading("5.2 Completeness: lexical vs semantic matching", level=2)
simple_box(
    "To grade extraction, we compare what EDIA found with the person's list of reasons. **Completeness** "
    "is the share of the person's reasons that EDIA also found.",
    "But how do you decide two sentences “match”? Strictly (same words) or loosely (same "
    "meaning)? Each way has a flaw, so the paper always reports **both**:",
    "- **Lexical (strict, same words):** under-credits correct answers that are reworded.",
    "- **Semantic (same meaning):** fairer to rewording, but more forgiving.",
    "The true value lies between the two, which is why the paper gives a **range** (33–58%).")
technical()
bullets([
    "**Lexical match:** token F1 ≥ 0.45 between an extracted element and a gold element.",
    "**Semantic match:** cosine similarity ≥ 0.55 between all-MiniLM-L6-v2 sentence embeddings "
    "(a small model that turns a sentence into a vector so that similar meanings land close together).",
    "**Why both matter, concretely:** EDIA's correct headline answer *“Kryptonite was chosen as the "
    "preferred energy source”* scores only 0.36 on the lexical matcher, below the threshold, so it "
    "counts as a miss. 14 of the 32 lexical misses are correct extractions like this. Lexical "
    "completeness is a **floor**.",
    "**Owner/approver** elements are excluded from completeness (they are rarely stated in these "
    "meetings) and the same rule is applied to every method.",
])

doc.add_heading("5.3 Precision and F1 for extraction", level=2)
simple_box(
    "Completeness only asks “did you find the right things?”. A system that writes down *everything* "
    "would score well on that. **Precision** asks the opposite: “of everything you wrote down, how much "
    "matches the person's list?” **F1** combines the two into one number, so neither “find "
    "little but be right” nor “write everything” can win.",
    "One caution: the person listed only the key reasons, so some things EDIA found that are real but "
    "were not on the list get counted as wrong. Our precision is therefore a **lower bound**: the real "
    "value is at least this high.")
technical()
bullets([
    "**Precision** = extracted elements matching some gold element at the completeness threshold / all "
    "extracted elements. **Recall** = completeness. **F1** = 2PR / (P + R).",
    "**Redundancy-penalised precision** counts each gold element at most once, exposing duplicates: EDIA "
    "extracts each chunk separately, so the same reason found in two chunks appears twice.",
])

doc.add_heading("5.4 Hallucination: the unsupported-quote rate", level=2)
simple_box(
    "EDIA must back every reason with an exact quote from the meeting. We check whether each quote "
    "really appears there. A quote that does not is **unsupported**: the AI may have invented it. This "
    "is our measure of hallucination, and it needs no human answers, so it can be measured on all 227 "
    "meetings.")
technical()
para("Each quote is searched for in its own source chunk after normalising case, punctuation and "
     "whitespace and ignoring transcription markers and filler words. The **unsupported-quote rate** is "
     "the fraction not found. A separate check flags “tidied” quotes: every word appears, in "
     "order, in the source, but some words were dropped. That is rewording, not invention.")

# 6 -------------------------------------------------------------------------------------------
doc.add_heading("6. Finding 1: spotting what is missing, and word overlap wins (paper Sec. V)", level=1)
simple_box(
    "The simplest detector, which just checks for shared words, beats the NLI “logic” model. The "
    "surprise is *why*: NLI keeps calling points “missing” even when the minutes contain them, "
    "sometimes word for word. Asking the AI model directly is no better than flipping a coin.")
technical()
table(["Detector", "Tier", "n", "Gap F1", "Bal. acc.", "Present F1"],
      [["NLI", "gold", "50", "0.861", "0.656", "0.476"],
       ["NLI", "synthetic", "1,667", "0.771", "0.629", "0.468"],
       ["Lexical", "gold", "50", "0.884", "0.583", "0.286"],
       ["Lexical", "synthetic", "1,667", "0.857", "**0.775**", "**0.715**"],
       ["Lexical", "corpus", "4,930", "0.970", "0.669", "0.457"],
       ["*Same 200-element sample:*", "", "", "", "", ""],
       ["NLI", "gold + synth.", "200", "–", "0.611", "0.432"],
       ["Lexical", "gold + synth.", "200", "–", "**0.728**", "**0.637**"],
       ["LLM judge", "gold + synth.", "200", "–", "0.492", "0.596"],
       ["Ensemble", "gold + synth.", "200", "–", "0.521", "0.624"]],
      [3.6, 2.6, 1.6, 1.8, 1.9, 2.0], font=9, align=["l", "l", "r", "r", "r", "r"])
caption("Paper Table II. Bal. acc. = balanced accuracy (0.5 = chance).")
bullets([
    "**On the balanced tier, lexical beats NLI on every honest metric** (balanced accuracy 0.775 vs "
    "0.629; present F1 0.715 vs 0.468), and the order holds when all four detectors share one sample.",
    "**Why NLI loses (accuracy by true label, gold + synthetic):** on true gaps NLI is fine (0.916; "
    "lexical 0.974), but it recognises paraphrased content only **36.6%** of the time and word-for-word "
    "copies only **29.2%** (lexical: 64.0% and 45.3%). Against a summary of about 100 words, almost every "
    "source sentence reads as “not implied”, so NLI is reacting to the summary being short, not "
    "doing useful entailment.",
    "**The LLM judge is at chance.** On 296 elements it called about 95% “present”, including "
    "obvious omissions: gap recall 0.047, balanced accuracy 0.495. We verified it received the right "
    "inputs, so this is a genuine yes-bias, a known weakness of models used as judges.",
    "**The ensemble inherits that bias** (0.516): calling an element present if *any* member says so "
    "lets the yes-saying judge override the others.",
])

# 7 -------------------------------------------------------------------------------------------
doc.add_heading("7. Finding 2: the extraction ceiling, and why it exists (paper Sec. VI)", level=1)
simple_box(
    "On the 5 hand-checked meetings, EDIA finds between a third and more than half of the reasons a "
    "person found, depending on how strictly a match is judged. When we read every miss by hand, most "
    "turned out to be reasoning that is **scattered** (the two halves of one reason said hundreds of "
    "turns apart) or **implicit** (never said outright). Not one miss was caused by how the meeting was "
    "cut into pieces, and not one was a made-up quote.")
technical()
table(["Meeting", "Lexical", "Semantic"],
      [["QSUM_001", "0.455", "0.636"], ["QSUM_002", "0.500", "0.875"], ["QSUM_003", "0.111", "0.222"],
       ["QSUM_004", "0.286", "0.643"], ["QSUM_005", "0.333", "0.500"],
       ["**Pooled**", "**0.333** (16/48)", "**0.583** (28/48)"]],
      [4.0, 4.0, 4.0], font=9.5, align=["l", "r", "r"])
caption("Paper Table III: completeness on the five gold meetings. The unsupported-quote rate here is 16.4%.")

doc.add_heading("Why the ceiling exists: the failure taxonomy", level=2)
table(["Failure category", "n", "Share", "What it means"],
      [["Multi-hop, implicit", "6", "30%", "The connection is never stated; it has to be inferred"],
       ["Multi-hop, scattered", "5", "25%", "The reasoning is split across turns far apart"],
       ["Paraphrase miss", "5", "25%", "Found, but worded so differently the scorer did not match it"],
       ["Out of scope", "4", "20%", "Stated plainly in one place, but the model did not extract it"],
       ["Chunk boundary", "0", "0%", "—"],
       ["Hallucination", "0", "0%", "—"]],
      [3.8, 1.0, 1.4, 10.4], font=9, align=["l", "r", "r", "l"])
caption("Paper Table V: all 20 misses remaining under semantic matching, categorised by reading the transcript.")
bullets([
    "**Multi-hop accounts for 11 of 20 misses (55%).** In one meeting, about **670 turns** separate the "
    "two halves of a single rationale.",
    "**The two zeros matter as much as the top category.** No chunk-boundary misses means a longer "
    "context window would not fix this; no fabricated quotes among the misses means the misses are "
    "omissions, not inventions.",
    "The taxonomy uses the *semantic* misses on purpose: the lexical misses include 14 correct answers "
    "the scorer under-credited, which would have invented a fake “paraphrase” problem.",
])

doc.add_heading("Grounding across all 227 meetings", level=2)
simple_box(
    "Completeness needs a person's answer list, which only the 5 gold meetings have. But the quote check "
    "needs none, so we ran it on all 227 meetings. Academic meetings looked much worse at first, but "
    "most of those quotes were just tidied: the AI dropped the “um”s, false starts and repeats "
    "common in those recordings. Counting only quotes with something genuinely wrong, all three domains "
    "fall between 3.6% and 7.6%, so EDIA behaves consistently across domains.")
technical()
table(["Domain", "Meetings", "Quotes", "Unsupported", "Excl. tidied"],
      [["Academic", "59", "6,341", "26.7%", "7.6%"], ["Committee", "36", "5,441", "5.1%", "3.6%"],
       ["Product", "132", "10,167", "12.0%", "5.1%"], ["**All**", "**227**", "**21,949**", "**14.6%**", "**5.4%**"]],
      [3.2, 2.2, 2.4, 2.8, 2.8], font=9.5, align=["l", "r", "r", "r", "r"])
caption("Paper Table IV. 3,532 chunks; 36 (1.0%) failed because the model got stuck repeating itself. "
        "72% of Academic's unsupported quotes are near-verbatim tidying.")

# 8 -------------------------------------------------------------------------------------------
doc.add_heading("8. Finding 3: other methods do not break the ceiling (paper Sec. VII)", level=1)
simple_box(
    "We tried three well-known alternatives on the same 5 meetings, with the same AI model wherever a "
    "model is needed:",
    "- **Think step by step (chain-of-thought):** the AI first writes notes, then fills in the form. It "
    "finds 2–3 more reasons out of 48, but writes 50% more, so once the extra output is counted it "
    "is **not** better overall.",
    "- **GraphRAG** (Microsoft's knowledge-graph method): much worse. It organises the meeting around "
    "people and things (“the Project Manager”) rather than around decisions and their reasons.",
    "- **Multi-step search (multi-hop retrieval):** search for the reason, then use it to search for the "
    "alternatives, and so on. Even when given the right decisions to start from, it finds very little: "
    "the first search already goes wrong and each step compounds the error.",
    "None beats EDIA, and each fails for a reason that points back to the same difficulty: the reasoning "
    "is scattered and implicit.")
technical()
table(["Approach", "Compl. (lex)", "Compl. (sem)", "Unsupp. quotes", "Core failure reason"],
      [["EDIA", "33.3%", "58.3%", "16.4%", "Multi-hop reasoning across chunks"],
       ["EDIA, same runtime †", "27.1%", "58.3%", "14.1%", "(the fair baseline for the † rows)"],
       ["EDIA + chain-of-thought †", "33.3%", "62.5%", "13.9%", "Better reasoning over content in view, not retrieval"],
       ["GraphRAG †", "4.2%", "14.6–18.8%", "–", "Entity graph built for structured text, not conversation"],
       ["Multi-hop retrieval", "0.0%", "12.8%", "0", "Chain unknown in advance; the first hop fails"]],
      [4.2, 1.9, 2.2, 2.2, 6.1], font=8.5, align=["l", "r", "r", "r", "l"])
caption("Paper Table VI. † = see “Why there is a same-runtime row” below.")
table(["Approach", "Elements", "Precision", "Recall", "F1"],
      [["EDIA", "398", "4.8 / 13.8", "33.3 / 58.3", "8.4 / **22.3**"],
       ["EDIA, same runtime †", "346", "4.9 / 13.6", "27.1 / 58.3", "8.3 / 22.0"],
       ["EDIA + chain-of-thought †", "438", "3.2 / 10.7", "33.3 / 62.5", "5.8 / 18.3"],
       ["GraphRAG †", "52", "3.8 / 13.5", "4.2 / 14.6", "4.0 / 14.0"],
       ["Multi-hop retrieval", "27", "0.0 / 40.7", "0.0 / 12.8", "0.0 / 19.5"]],
      [4.6, 1.8, 2.8, 2.8, 2.6], font=9, align=["l", "r", "r", "r", "r"])
caption("Paper Table VII: precision, recall and F1, each as lexical / semantic (%). Precision is a lower bound.")
bullets([
    "**Chain-of-thought.** Implemented properly as two calls per chunk: one free-text call that writes "
    "reasoning notes, then the normal form-filling call that reads them. (A single form-filling call "
    "cannot think aloud: it is forced to output the form directly.) Against the fair baseline it raises "
    "recall by 6.2 points lexical and 4.2 semantic (3 and 2 of 48 elements) with unchanged hallucination, "
    "but emits 50% more records (114 vs 76); its semantic precision falls from 13.6% to 10.7% and its F1 "
    "from **22.0% to 18.3%**. The recall gain does not survive.",
    "**GraphRAG.** Run end-to-end with Microsoft's default configuration and its local search. Its prose "
    "answer is converted into our form by one call told to restate only; items not supported by the "
    "answer are dropped (the converter was seen inventing an alternative). 4.2% lexical, 14.6% semantic, "
    "or 18.8% semantic without that filter. Hence the range.",
    "**Multi-hop retrieval.** A three-hop chain (rationale, then alternatives, then objections) using "
    "keyword-style search (TF-IDF) over transcript turns, handed the true decisions. It is scored **only on "
    "what it retrieves**: the decision text it is handed often resembles the true rationale, and letting "
    "that match would credit the method with its own input (that would give 12.8% / 28.2%). Lexical is 0% "
    "because a retrieved turn is long and the true element short, so their word overlap rarely clears "
    "the threshold.",
    "**Reading Table VII.** EDIA has the best semantic F1 (22.3%). Precision is low for every method that "
    "writes its own answers, because the gold list has only 48 key elements while EDIA writes 398, many "
    "of them the same reason found again in another chunk (counting each gold element once lowers "
    "EDIA's precision to 7.0%) and many plausible reasons the annotator did not list. So compare "
    "precision *between* methods; do not read it as an error rate. Multi-hop's high precision comes from "
    "writing very little (27 elements copied from the transcript), not from finding more.",
])

doc.add_heading("Why there is a “same runtime” row", level=2)
simple_box(
    "The program that runs the AI model was updated partway through the project. Same model, same "
    "settings, yet a few outputs changed slightly. To keep comparisons fair, we re-ran plain EDIA on the "
    "updated program and compare the newer experiments (marked †) against that re-run, not against "
    "the original. Comparing like with like is the only fair test.")
technical()
para("Replaying the plain extractor on the newer inference runtime, with identical weights and "
     "deterministic settings, moved lexical completeness by −6.2 points and left semantic "
     "completeness unchanged. Even within one session, repeating the LLM judge on the same 296 elements "
     "gave a different confidence for 7 of them (3 verdicts changed). The paper therefore recommends "
     "reporting the inference-engine version alongside the model, and not treating temperature-0 "
     "decoding as perfectly reproducible.")

# 9 -------------------------------------------------------------------------------------------
doc.add_heading("9. What it all means (paper Sec. VIII–IX)", level=1)
simple_box(
    "The limit is not a bug in one method. It comes from how meetings work: the reasons behind a "
    "decision are spread across the conversation and often left unsaid. Methods that look at one piece "
    "at a time, or search by similar words, cannot connect them. Breaking the ceiling needs a system "
    "that follows a decision through the whole meeting: linking a later comment back to the proposal it "
    "responds to, *before* trying to extract anything.")
technical()
bullets([
    "**Ruled out as fixes:** a longer context window (no chunk-boundary misses), better reasoning at "
    "answer time (chain-of-thought's small recall gain is offset by lower precision), entity graphs "
    "(GraphRAG) and chained similarity search (multi-hop retrieval).",
    "**Likely direction:** discourse-level tracking of each decision across the whole meeting before "
    "extraction, rather than retrieval by surface similarity.",
    "**Two reporting practices the paper recommends generally:** report balanced accuracy and "
    "present-class F1 alongside gap F1 when gaps dominate, and report both lexical and semantic "
    "completeness. Both changed this paper's conclusions.",
])

# 10 ------------------------------------------------------------------------------------------
doc.add_heading("10. Limitations: what the paper does not claim", level=1)
simple_box(
    "Every study has limits, and saying them plainly makes the paper stronger. Ours: the most important "
    "number (completeness) comes from only 5 meetings of one kind; the large corpus is labelled by an AI; "
    "the human check is small; the made-up meetings are cleaner than real ones; and we only tried small, "
    "free models.")
technical()
bullets([
    "**Completeness on five Product meetings only.** Across all three domains we measured grounding "
    "(hallucination), not completeness, so we cannot claim the 33–58% ceiling holds everywhere.",
    "**Silver corpus labels.** Produced by NLI, which is weak on present content, so they overstate the gap rate.",
    "**Small human spot-check.** One annotator, 52 elements from three pairs, no Product pair; no second "
    "annotator, so no inter-annotator agreement (Cohen's kappa).",
    "**Synthetic meetings** are model-written and cleaner than real speech.",
    "**QMSum summaries stand in for filed decision records**; real minutes may differ.",
    "**Only 7B-scale open models**, run locally; larger models may behave differently (future work).",
])

# 11 ------------------------------------------------------------------------------------------
doc.add_heading("11. Questions you may be asked, and how to answer them", level=1)
qa("Isn't finding only 33–58% of the reasoning a failure?",
   "It is a feasibility study: the goal was to find out *whether* and *where* the task works, not to "
   "sell a product. Measuring the ceiling and explaining why it exists (multi-hop reasoning, not "
   "chunking or hallucination) is the result.")
qa("Why give a range (33–58%) instead of one number?",
   "Because there are two fair ways to decide whether two sentences match, and each has a known flaw. "
   "The strict one misses correct reworded answers (14 of its 32 misses are correct), and the loose one "
   "is more forgiving. Reporting only one would hide that.")
qa("Why do you use balanced accuracy instead of accuracy or F1?",
   "Because 93% of corpus elements are gaps. Always guessing “gap” would score about 93% "
   "accuracy and an F1 of about 0.96 while doing nothing. Balanced accuracy gives that lazy detector 0.5.")
qa("NLI is a standard tool for checking summaries. Why did it lose to word overlap?",
   "NLI is normally used the other way round: checking that a *summary* is supported by its *source*. "
   "Gap detection asks whether a source sentence survives in a much shorter summary. In that direction "
   "NLI calls almost everything “not implied”, even word-for-word copies (29.2% accuracy).")
qa("Would a bigger AI model fix the LLM judge or the extractor?",
   "Possibly, and we say so as a limitation: we only tested 7B open models. But the judge's problem is "
   "a systematic yes-bias, and the extraction misses are about scattered and implicit reasoning, which "
   "the alternatives we tested did not fix either.")
qa("Why is completeness measured on only 5 meetings?",
   "Completeness needs a person to list every reason in the meeting by hand, which takes hours per "
   "meeting. Across all 227 meetings we measured what needs no human answers, the quote check, and the "
   "paper states this limitation explicitly.")
qa("EDIA's precision is 13.8%. Does that mean it is mostly wrong?",
   "No. The answer list has only the 48 key reasons; EDIA writes 398 elements, many of them the same "
   "reason found again in another chunk, and many real reasons that are not on the short list. Precision "
   "is a lower bound for comparing methods, and on F1, which balances both, EDIA is the best method.")
qa("Did chain-of-thought help?",
   "It found slightly more (+4.2 points semantic recall) but wrote 50% more, so its precision and F1 "
   "dropped (F1 22.0% to 18.3%). Overall, no.")
qa("Why does multi-hop retrieval score 0% lexical?",
   "It returns whole transcript turns, which are long, while the true reasons are short phrases, so "
   "word overlap rarely clears the threshold. The semantic score (12.8%) is the informative one. We "
   "also score it only on what it retrieves, not on the decisions we handed it.")
qa("Does EDIA make things up?",
   "Some quotes cannot be found in the transcript: 16.4% on the gold meetings, 14.6% across the corpus. "
   "Most of those are tidied quotes with filler words removed; counting only genuinely wrong quotes "
   "it is 5.4% across the corpus. None of the 20 analysed misses was caused by an invented quote.")
qa("If an AI labelled the 227 corpus meetings, can we trust those numbers?",
   "Only as upper bounds on the gap rate, which is how the paper uses them. The human spot-check shows "
   "the labels never missed a real gap but called about half of present content “missing”. We "
   "never score NLI against its own labels.")
qa("Why use made-up meetings at all?",
   "They are the only data where the right answers are known by construction and the gaps are close to "
   "50/50, which is what a fair detector comparison needs. Their limitation, being cleaner than real "
   "speech, is stated in the paper.")
qa("What would you do next?",
   "Track each decision's discussion across the whole meeting before extracting (link later turns back "
   "to the proposal they respond to); hand-label a few Academic and Committee meetings to test the "
   "ceiling across domains; add a second human annotator; and test larger models.")

# 12 ------------------------------------------------------------------------------------------
doc.add_heading("12. Map of the paper", level=1)
table(["In the paper", "What it shows", "Explained here"],
      [["Abstract, Sec. I", "The problem, the four contributions, the one-word answer (“partially”)", "§§1–2"],
       ["Sec. II", "Related work: QMSum, NLI for faithfulness, LLM judges, RAG/GraphRAG, chain-of-thought", "§3, glossary"],
       ["Sec. III, Fig. 1", "The EDIA pipeline and how it is evaluated", "§3, §5"],
       ["Sec. IV, Table I", "The edia_db benchmark and the human spot-check", "§4"],
       ["Sec. V, Table II", "Gap detection: lexical beats NLI; LLM judge at chance", "§6"],
       ["Sec. VI, Tables III–V", "Completeness, corpus grounding, failure taxonomy", "§7"],
       ["Sec. VII, Tables VI–VII", "Comparison with CoT, GraphRAG, multi-hop; precision/F1", "§8"],
       ["Sec. VIII–IX", "Discussion, limitations, conclusion", "§§9–10"]],
      [3.8, 9.4, 3.4], font=9)

# 13 ------------------------------------------------------------------------------------------
doc.add_heading("13. Glossary", level=1)
table(["Term", "In simple terms", "Technically"],
      [["Decision-context gap", "A reason said in the meeting but left out of the minutes", "Reasoning element absent from the filed document (label “omitted”)"],
       ["Reasoning element", "One piece of “why”: an option, a reason, a worry, a condition", "Typed unit: decision / alternative / rationale / condition / dissent / owner / approver"],
       ["Gold / silver / constructed labels", "Checked by a person / guessed by a model / known by design", "Human-annotated / NLI-produced / fixed before text generation"],
       ["Transcript, chunk", "The written record of the meeting; a bite-sized piece of it", "Speaker-turn text; ≤ 1,200 tokens with 2-turn overlap"],
       ["Token", "Roughly a word or part of a word", "Unit the model reads and writes"],
       ["Mistral 7B", "The free AI model EDIA uses", "Open 7-billion-parameter language model"],
       ["NLI (entailment)", "A model that judges whether one text implies another", "roberta-large-mnli, P(entailment | document, element)"],
       ["LLM judge", "Asking the AI model “is this in there?”", "Mistral 7B prompted for a verdict and confidence"],
       ["Threshold", "The cut-off score for saying yes", "0.55 for detectors; 0.45 / 0.55 for lexical / semantic matching"],
       ["Base rate", "How common the “missing” answer is", "Prior probability of the gap class (0.93 in the corpus)"],
       ["Precision / recall / F1", "Right among what you said / found among what exists / both combined", "TP/(TP+FP); TP/(TP+FN); harmonic mean"],
       ["Balanced accuracy", "A fair score where blind guessing gets 50%", "Mean of per-class recall"],
       ["Completeness", "Share of the person's reasons the system found", "Matched gold elements / all gold elements (= recall)"],
       ["Lexical vs semantic matching", "Same words vs same meaning", "Token F1 ≥ 0.45 vs MiniLM cosine ≥ 0.55"],
       ["Unsupported quote", "Evidence the AI gave that is not in the transcript", "Quote not found in its own chunk after normalisation"],
       ["Deterministic decoding", "The AI always picks its single most likely next word", "Temperature 0, top-k 1, fixed seed"],
       ["Chain-of-thought (CoT)", "Asking the AI to think step by step first", "Free-text notes pass, then constrained extraction"],
       ["GraphRAG", "Microsoft's method: build a map of people and things, then search it", "Entity graph + community summaries; local search"],
       ["TF-IDF", "Search that favours rare, distinctive words", "Term-frequency × inverse-document-frequency vectors, cosine similarity"],
       ["Multi-hop retrieval", "Search, then use the result to search again", "Chained retrieval: rationale → alternatives → objections"],
       ["Multi-hop reasoning", "Connecting clues spread far apart", "Evidence for one element spans distant, non-adjacent turns"],
       ["Inference runtime", "The program that runs the AI model", "Local model server; its version can change outputs slightly"]],
      [3.6, 6.0, 7.0], font=8.5)

doc.core_properties.title = "EDIA, explained"
doc.core_properties.author = "Jaitra Shah; Manan Mehta; Soham Dwivedi; Aayusman Pati"
doc.save(OUT)
print(f"wrote {OUT}")
