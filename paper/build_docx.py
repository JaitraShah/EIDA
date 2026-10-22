"""Build paper/edia_ieee.docx from paper/edia_ieee.tex + paper/refs.bib (IEEE conference layout).

    .venv\\Scripts\\python.exe paper/make_figure.py     # Fig. 1 PNG (shared with the LaTeX build)
    .venv\\Scripts\\python.exe paper/build_docx.py      # -> paper/edia_ieee.docx

The .tex is the single source: edit it, then re-run this script. This is a converter
for the constructs edia_ieee.tex actually uses (sections, subsections, paragraphs,
itemize, table/table*, figure, \\textbf/\\emph/\\textit/\\texttt/\\url/\\footnote/\\TODO,
\\cite, \\ref, simple inline math) -- not a general LaTeX engine.

Layout follows the IEEE conference template (Template (1).docx): US Letter, margins
0.75/1/0.62/0.62 in, two columns, Times New Roman 10 pt. Paragraphs carry the template's own
style names (paper title, Author, Abstract, Keywords, Heading 1/2, Body Text, bullet list,
figure caption, table head, table copy, references), defined in make_styles() with the
template's properties, so Word's Styles pane shows the same vocabulary. The four authors sit
in a four-column section, as the template prescribes for fewer than six authors.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEX = HERE / "edia_ieee.tex"
BIB = HERE / "refs.bib"
OUT = HERE / "edia_ieee.docx"
FONT = "Times New Roman"
COL_W = 3.45      # usable single-column width (in)
FULL_W = 7.1      # usable full-page width (in)
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
ACCENTS = {r"{\'s}": "ś", r"{\'n}": "ń", r"\'e": "é", r"\'a": "á", r'\"o': "ö", r'\"u': "ü"}


# ----------------------------------------------------------------------------- tex helpers
def strip_comments(tex: str) -> str:
    return "\n".join(re.sub(r"(?<!\\)%.*$", "", line) for line in tex.splitlines())


def braced(s: str, i: int) -> tuple[str, int]:
    """s[i] == '{'. Return (content, index after the closing brace)."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while j < len(s):
        if s[j] == "\\":
            j += 2
            continue
        if s[j] == "{":
            depth += 1
        elif s[j] == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    raise ValueError("unbalanced braces near: " + s[i:i + 60])


def command_arg(s: str, name: str) -> str | None:
    """First mandatory {argument} of \\name, skipping an optional [..] argument."""
    m = re.search(r"\\" + re.escape(name) + r"\s*(?:\[[^\]]*\])?\s*\{", s)
    return braced(s, m.end() - 1)[0] if m else None


def env(s: str, name: str) -> str | None:
    m = re.search(r"\\begin\{" + re.escape(name) + r"\}(.*?)\\end\{" + re.escape(name) + r"\}", s, re.S)
    return m.group(1) if m else None


# ----------------------------------------------------------------------------- inline -> runs
def typography(t: str) -> str:
    t = t.replace("---", "\u2014").replace("--", "\u2013")
    t = t.replace("``", "\u201c").replace("''", "\u201d").replace("`", "\u2018").replace("'", "\u2019")
    return t


def math_runs(m: str, fmt: dict) -> list:
    for a, b in ((r"\geq", "≥"), (r"\leq", "≤"), (r"\times", "×"), (r"\dagger", "†"), ("-", "−")):
        m = m.replace(a, b)
    runs, i = [], 0
    while i < len(m):
        ch = m[i]
        if ch == "^":
            nxt = m[i + 1] if i + 1 < len(m) else ""
            if nxt == "{":
                sup, i = braced(m, i + 1)
            else:
                sup, i = nxt, i + 2
            runs.append((sup, {**fmt, "sup": True}))
            continue
        j = i
        if ch.isalpha():
            while j < len(m) and m[j].isalpha():
                j += 1
            runs.append((m[i:j], {**fmt, "i": True}))
        else:
            while j < len(m) and not m[j].isalpha() and m[j] != "^":
                j += 1
            runs.append((m[i:j], dict(fmt)))
        i = j
    return runs


SYMBOLS = {"%": "%", "&": "&", "_": "_", "#": "#", ",": "\u2009", " ": " ", "$": "$", "{": "{", "}": "}"}
STYLE_CMDS = {"textbf": {"b": True}, "emph": {"i": True}, "textit": {"i": True},
              "texttt": {"mono": True}, "textsc": {"sc": True}, "url": {}}


def inline(s: str, fmt: dict | None = None) -> list:
    """LaTeX inline text -> list of (text, fmt) runs."""
    fmt = fmt or {}
    runs, buf, i = [], [], 0

    def flush():
        if buf:
            runs.append((typography("".join(buf)), dict(fmt)))
            buf.clear()

    while i < len(s):
        ch = s[i]
        if ch == "\\":
            m = re.match(r"\\([A-Za-z]+)\*?", s[i:])
            if not m:                                   # \%, \&, \_, \,, "\ ", \\ ...
                sym = s[i + 1] if i + 1 < len(s) else ""
                if sym == "\\" or sym.isspace():
                    buf.append(" ")
                else:
                    buf.append(SYMBOLS.get(sym, sym))
                i += 2
                continue
            name, i = m.group(1), i + m.end()
            while i < len(s) and s[i] == " " and name not in STYLE_CMDS and name not in ("TODO", "footnote"):
                i += 1  # control words eat following spaces
            if name in STYLE_CMDS or name in ("TODO", "footnote"):
                arg, i = braced(s, i)
                flush()
                if name == "TODO":
                    red = {**fmt, "color": "C00000"}
                    runs += [("[TODO: ", red)] + inline(arg, red) + [("]", red)]
                elif name == "footnote":
                    runs += [(", ", dict(fmt))] + inline(arg, fmt)
                elif name == "url":
                    runs.append((arg, dict(fmt)))
                else:
                    runs += inline(arg, {**fmt, **STYLE_CMDS[name]})
            elif name == "ldots":
                buf.append("\u2026")
            # any other control word is layout-only here (\centering, \small, ...) and dropped
            continue
        if ch == "$":
            j = s.index("$", i + 1)
            flush()
            runs += math_runs(s[i + 1:j], fmt)
            i = j + 1
            continue
        if ch in "{}":
            i += 1
            continue
        if ch == "~":
            buf.append("\u00a0")
        elif ch in "\n\t":
            buf.append(" ")
        else:
            buf.append(ch)
        i += 1
    flush()
    # collapse whitespace across run boundaries
    out, prev_space = [], True
    for text, f in runs:
        text = re.sub(r"[ ]{2,}", " ", text)
        if prev_space:
            text = text.lstrip(" ")
        if text:
            out.append((text, f))
            prev_space = text.endswith(" ")
    if out:
        out[-1] = (out[-1][0].rstrip(" "), out[-1][1])
    return out


# ----------------------------------------------------------------------------- docx helpers
def set_run(run, fmt: dict, size: float):
    if size is not None:
        run.font.size = Pt(size)
    run.bold = fmt.get("b") or None
    run.italic = fmt.get("i") or None
    if fmt.get("sc"):
        run.font.small_caps = True
    if fmt.get("sup"):
        run.font.superscript = True
    if fmt.get("mono"):
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    if fmt.get("color"):
        run.font.color.rgb = RGBColor.from_string(fmt["color"])


def add_runs(par, runs, size=None, base: dict | None = None):
    for text, fmt in runs:
        set_run(par.add_run(text), {**(base or {}), **fmt}, size)


def para(doc, runs=None, style="Body Text", align=None, indent=None, before=None, after=None,
         keep_next=False, size=None):
    """Paragraph in a template style; only the arguments given override the style."""
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if indent is False:
        pf.first_line_indent = Pt(0)
    if keep_next:
        pf.keep_with_next = True
    if runs:
        add_runs(p, runs, size)
    return p


def spacer(doc, pts: float):
    """Empty paragraph `pts` tall."""
    p = para(doc, style="Normal", before=0, after=0)
    p.paragraph_format.line_spacing = Pt(pts)
    return p


def _style(doc, name, base="Normal", size=None, bold=None, italic=None, small_caps=None,
           align=None, before=None, after=None, line=None, first=None, left=None, spacing=None):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles[base]
    st.quick_style = True
    f, pf = st.font, st.paragraph_format
    if size is not None:
        f.size = Pt(size)
    f.bold, f.italic, f.small_caps = bold, italic, small_caps
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if first is not None:
        pf.first_line_indent = Pt(first)
    if left is not None:
        pf.left_indent = Pt(left)
    if spacing is not None:                                   # character spacing, 1/20 pt
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        st.element.get_or_add_rPr().append(sp)
    return st


def make_styles(doc):
    """The template's paragraph styles (names and properties from Template (1).docx)."""
    J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.alignment = C                     # template Normal is centred
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    for name in ("Heading 1", "Heading 2", "Body Text"):      # drop Word's themed versions
        el = styles[name].element
        el.getparent().remove(el)
    _style(doc, "paper title", size=24, align=C, after=6)
    _style(doc, "Author", size=11, align=C, before=18, after=2)
    _style(doc, "Abstract", size=9, bold=True, align=J, after=10, first=13.6)
    _style(doc, "Keywords", base="Abstract", italic=True, after=6, first=13.7)
    heads = []
    for name, lvl, kw in (("Heading 1", 0, dict(small_caps=True, align=C, before=8, after=4, first=0)),
                          ("Heading 2", 1, dict(italic=True, align=L, before=6, after=3))):
        st = _style(doc, name, **kw)
        st.paragraph_format.keep_with_next = st.paragraph_format.keep_together = True
        ol = OxmlElement("w:outlineLvl")
        ol.set(qn("w:val"), str(lvl))
        st.element.get_or_add_pPr().append(ol)
        heads.append(st)
    body = _style(doc, "Body Text", align=J, after=6, line=0.95, first=14.4, spacing=-1)
    for st in heads:
        st.next_paragraph_style = body
    _style(doc, "bullet list", base="List Bullet", align=J, after=6, line=0.95, left=28.8,
           first=-14.4, spacing=-1)
    _style(doc, "figure caption", size=8, align=J, before=4, after=10, first=0)
    _style(doc, "table head", size=8, small_caps=True, align=C, before=12, after=6, line=0.9)
    _style(doc, "table copy", size=8, align=J)
    ref = _style(doc, "references", size=8, align=J, after=2.5, left=21.6, first=-21.6)
    ref.paragraph_format.line_spacing = Pt(9)                 # exact 9 pt
    ref.paragraph_format.tab_stops.add_tab_stop(Pt(21.6))


def set_columns(section, n: int, space_in: float = 0.25):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_in * 1440)))


def page_setup(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Pt(54), Pt(72)
    section.left_margin = section.right_margin = Pt(44.65)


# WordprocessingML fixes the child order of these property blocks; Word treats an
# out-of-order child as corruption (it opens a hidden "repair?" prompt under automation).
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def insert_ordered(parent, child, order):
    """Insert child at its schema position, replacing any existing element of the same name."""
    name = child.tag.split("}")[1]
    for old in parent.findall(qn("w:" + name)):
        parent.remove(old)
    rank = order.index(name)
    for sib in parent:
        sname = sib.tag.split("}")[1]
        if sname in order and order.index(sname) > rank:
            sib.addprevious(child)
            return
    parent.append(child)


def cell_borders(cell, **edges):
    """edges: top/bottom=<size in eighths of a point>. Idempotent (merged cells repeat)."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    insert_ordered(tcPr, borders, TCPR_ORDER)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(edges[edge]))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


def no_table_borders(table):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        b.append(el)
    insert_ordered(table._tbl.tblPr, b, TBLPR_ORDER)


def cell_margins(table, left=0.05, right=0.05):
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 0.0), ("left", left), ("bottom", 0.0), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(v * 1440)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    insert_ordered(table._tbl.tblPr, mar, TBLPR_ORDER)


# ----------------------------------------------------------------------------- bibliography
def parse_bib(text: str) -> dict:
    entries = {}
    for m in re.finditer(r"@(\w+)\s*\{\s*([^,\s]+)\s*,", text):
        kind, key = m.group(1).lower(), m.group(2)
        i, fields = m.end(), {}
        while True:
            fm = re.compile(r"\s*(\w+)\s*=\s*").match(text, i)
            if not fm:
                break
            name, j = fm.group(1).lower(), fm.end()
            if text[j] == "{":
                val, j = braced(text, j)
            else:
                vm = re.compile(r"[^,}\n]+").match(text, j)
                val, j = vm.group(0).strip(), vm.end()
            fields[name] = val
            cm = re.compile(r"\s*,?").match(text, j)
            i = cm.end()
            if i < len(text) and text[i] == "}":
                break
        entries[key] = {"kind": kind, **fields}
    return entries


def bib_text(s: str) -> str:
    for a, b in ACCENTS.items():
        s = s.replace(a, b)
    return typography(re.sub(r"[{}]", "", s).replace("\\&", "&"))


def ieee_authors(raw: str) -> str:
    names = [n.strip() for n in bib_text(raw).split(" and ")]
    others = "others" in names
    names = [n for n in names if n != "others"]

    def initials(n):
        parts = n.split()
        return " ".join(p[0] + "." for p in parts[:-1]) + " " + parts[-1] if len(parts) > 1 else n

    names = [initials(n) for n in names]
    if others or len(names) > 6:
        return names[0] + " et al."
    if len(names) == 1:
        return names[0]
    if len(names) == 2:
        return f"{names[0]} and {names[1]}"
    return ", ".join(names[:-1]) + ", and " + names[-1]


def ieee_reference(e: dict) -> list:
    """-> runs for one reference entry."""
    runs = [(ieee_authors(e["author"]) + ", \u201c" + bib_text(e["title"]) + ",\u201d ", {})]
    pages = f", pp. {bib_text(e['pages'])}" if "pages" in e else ""
    if e["kind"] == "inproceedings":
        runs += [("in ", {}), (bib_text(e["booktitle"]), {"i": True})]
        if "publisher" in e:
            runs.append((", " + bib_text(e["publisher"]), {}))
        runs.append((f", {e['year']}{pages}.", {}))
    else:
        runs.append((bib_text(e["journal"]), {"i": True}))
        vol = f", vol. {e['volume']}" if "volume" in e else ""
        runs.append((f"{vol}{pages}, {e['year']}.", {}))
    return runs


# ----------------------------------------------------------------------------- tables / figure
def parse_tabular(body: str):
    m = re.search(r"\\begin\{tabular\}\{(.*?)\}\s*\n(.*?)\\end\{tabular\}", body, re.S)
    spec = re.sub(r"@\{\}", "", m.group(1))
    aligns = [c for c in spec if c in "lrc"]
    rows, rules_after, pending_rule = [], set(), False
    for raw in re.split(r"\\\\", m.group(2)):
        lines = [l.strip() for l in raw.strip().splitlines()]
        for l in lines:
            if l in ("\\midrule", "\\toprule", "\\bottomrule"):
                if l == "\\midrule" and rows:
                    rules_after.add(len(rows) - 1)
        content = " ".join(l for l in lines if l not in ("\\midrule", "\\toprule", "\\bottomrule")).strip()
        if not content:
            continue
        if content.startswith("\\multicolumn"):
            i = len("\\multicolumn")
            _, i = braced(content, i)          # column count
            _, i = braced(content, i)          # alignment spec, may nest braces: {@{}l}
            text, _ = braced(content, i)
            rows.append(("span", text))
        else:
            rows.append(("cells", [c.strip() for c in re.split(r"(?<!\\)&", content)]))
    return aligns, rows, rules_after


def render_table(doc, block: str, labels: dict, width_in: float):
    caption = command_arg(block, "caption")
    label = command_arg(block, "label")
    aligns, rows, rules_after = parse_tabular(block)
    ncol = len(aligns)
    # caption above: "TABLE I" + caption text, small caps, centred, 8 pt
    para(doc, [(f"TABLE {labels[label]}", {})], style="table head", before=8, after=0, keep_next=True)
    para(doc, inline(caption), style="table head", before=0, after=4, keep_next=True)
    # column widths proportional to longest cell text
    plain = lambda t: "".join(x for x, _ in inline(t))
    longest = [max((len(plain(r[1][c])) for r in rows if r[0] == "cells"), default=4) for c in range(ncol)]
    weights = [min(max(L, 3), 60) + 3 for L in longest]      # +3 chars: cell margins
    widths = [width_in * w / sum(weights) for w in weights]
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_table_borders(table)
    cell_margins(table)
    for ri, (kind, content) in enumerate(rows):
        row = table.rows[ri]
        if kind == "span":
            merged = row.cells[0].merge(row.cells[-1])
            cells_content = [(merged, content, "l")]
        else:
            cells_content = [(row.cells[c], content[c], aligns[c]) for c in range(ncol)]
        for cell, text, al in cells_content:
            p = cell.paragraphs[0]
            p.style = doc.styles["table copy"]
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                                            "c": WD_ALIGN_PARAGRAPH.CENTER}[al]
            p.paragraph_format.keep_with_next = ri < len(rows) - 1 or None   # never split a table
            add_runs(p, inline(text))
        for c, cell in enumerate(row.cells):
            cell.width = Inches(widths[c])
            edges = {}
            if ri == 0:
                edges["top"] = 8                     # \toprule
            if ri in rules_after:
                edges["bottom"] = 4                  # \midrule
            if ri == len(rows) - 1:
                edges["bottom"] = 8                  # \bottomrule
            cell_borders(cell, **edges)
    spacer(doc, 6)                                   # small gap after the table


def render_figure(doc, block: str, labels: dict):
    img = command_arg(block, "includegraphics")
    caption = command_arg(block, "caption")
    label = command_arg(block, "label")
    p = para(doc, style="Normal", before=6, keep_next=True)
    p.add_run().add_picture(str(HERE / img), width=Inches(COL_W))
    para(doc, [(f"Fig. {labels[label]}. ", {})] + inline(caption), style="figure caption", after=8)


# ----------------------------------------------------------------------------- main
def main() -> int:
    tex = strip_comments(TEX.read_text(encoding="utf-8"))
    # expand argument-less \newcommand macros (e.g. \EDIAfull) as LaTeX would
    for name, val in re.findall(r"\\newcommand\{\\(\w+)\}\{([^{}]*)\}", tex):
        tex = re.sub(r"\\" + name + r"(?![A-Za-z])(\{\})?", lambda _m, v=val: v, tex)
    bib = parse_bib(BIB.read_text(encoding="utf-8"))
    body = tex.split("\\maketitle", 1)[1].split("\\bibliographystyle", 1)[0]
    for front in ("abstract", "IEEEkeywords"):
        body = re.sub(r"\\begin\{" + front + r"\}.*?\\end\{" + front + r"\}", "", body, flags=re.S)

    # numbering: citations in order of first appearance; tables Roman, figures Arabic
    order = []
    for grp in re.findall(r"\\cite\{([^}]+)\}", body):
        for k in grp.split(","):
            if k.strip() not in order:
                order.append(k.strip())
    missing = [k for k in order if k not in bib]
    assert not missing, f"cite keys missing from refs.bib: {missing}"
    cite_no = {k: i + 1 for i, k in enumerate(order)}
    labels, nt, nf = {}, 0, 0
    for m in re.finditer(r"\\begin\{(table\*?|figure)\}(.*?)\\end\{\1\}", body, re.S):
        lab = command_arg(m.group(2), "label")
        if m.group(1).startswith("table"):
            labels[lab] = ROMAN[nt]; nt += 1
        else:
            nf += 1; labels[lab] = str(nf)

    # section labels: a \label directly after a numbered \section gets its Roman number
    n_sec = 0
    for m in re.finditer(r"\\section(\*?)\s*\{[^}]*\}(\s*\\label\{([^}]+)\})?", body):
        if m.group(1):
            continue
        if m.group(3):
            labels[m.group(3)] = ROMAN[n_sec]
        n_sec += 1
    body = re.sub(r"(\\section\s*\{[^}]*\})\s*\\label\{[^}]+\}", r"\1", body)

    def resolve(s: str) -> str:
        s = re.sub(r"\\cite\{([^}]+)\}",
                   lambda m: ", ".join(f"[{cite_no[k.strip()]}]" for k in m.group(1).split(",")), s)
        return re.sub(r"\\ref\{([^}]+)\}", lambda m: labels[m.group(1)], s)

    doc = Document()
    make_styles(doc)
    page_setup(doc.sections[0])
    set_columns(doc.sections[0], 1)

    # --- title + authors (single column)
    title = command_arg(tex, "title")
    p = para(doc, style="paper title", after=10)
    add_runs(p, inline(re.sub(r"\\\\", " ", title)))       # Word wraps; no forced break
    author_src = command_arg(tex, "author")
    blocks = []
    for m in re.finditer(r"\\IEEEauthorblockN\{", author_src):
        name, j = braced(author_src, m.end() - 1)
        k = author_src.index("\\IEEEauthorblockA{", j)
        aff, _ = braced(author_src, k + len("\\IEEEauthorblockA"))
        blocks.append((name, [l for l in re.split(r"\\\\", aff) if l.strip()]))

    # --- authors: one column each, as the template does for fewer than six authors
    auth_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(auth_sec, len(blocks), 0.15)
    for n, (name, aff) in enumerate(blocks):
        p = para(doc, style="Author", before=6)
        if n:
            p.add_run().add_break(WD_BREAK.COLUMN)
        add_runs(p, inline(name))
        for line in aff:
            para(doc, inline(line), style="Author", before=0, after=0, size=9)

    # --- two-column body
    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body_sec, 2)

    abstract = env(tex, "abstract")
    p = para(doc, style="Abstract", before=6)
    add_runs(p, [("Abstract", {"i": True}), ("\u2014", {})] + inline(resolve(abstract)))
    kw = env(tex, "IEEEkeywords")
    p = para(doc, style="Keywords")
    add_runs(p, [("Keywords\u2014", {})] + inline(kw))

    sec_no, sub_no = 0, 0
    for chunk in re.split(r"\n\s*\n", body):
        chunk = chunk.strip()
        while chunk.startswith("\\section") or chunk.startswith("\\subsection"):
            m = re.match(r"\\(sub)?section(\*?)\s*\{", chunk)
            head, end = braced(chunk, m.end() - 1)
            if m.group(1):                                   # subsection
                label = f"{chr(ord('A') + sub_no)}. "
                sub_no += 1
                para(doc, [(label, {})] + inline(resolve(head)), style="Heading 2")
            else:
                if m.group(2):                               # unnumbered (\section*)
                    runs = inline(head)
                else:
                    runs = [(ROMAN[sec_no] + ". ", {})] + inline(resolve(head))
                    sec_no += 1
                    sub_no = 0
                para(doc, runs, style="Heading 1")
            chunk = chunk[end:].strip()
        # a chunk may mix paragraph text and environments (e.g. "...:" directly followed by itemize)
        while chunk:
            m_env = re.search(r"\\begin\{(table\*?|figure|itemize)\}", chunk)
            if not m_env:
                para(doc, inline(resolve(chunk)))
                break
            if m_env.start() > 0:
                para(doc, inline(resolve(chunk[:m_env.start()].strip())))
            kind = m_env.group(1)
            end_tag = "\\end{" + kind + "}"
            stop = chunk.index(end_tag, m_env.start()) + len(end_tag)
            block, chunk = chunk[m_env.start():stop], chunk[stop:].strip()
            if kind == "itemize":
                for item in re.split(r"\\item\s+", env(block, "itemize"))[1:]:
                    p = doc.add_paragraph(style="bullet list")
                    add_runs(p, inline(resolve(item.strip())))
            elif kind == "figure":
                render_figure(doc, resolve(block), labels)
            elif kind == "table":
                render_table(doc, resolve(block), labels, COL_W)
            else:                                            # table*: full width
                one = doc.add_section(WD_SECTION.CONTINUOUS)
                set_columns(one, 1)
                render_table(doc, resolve(block), labels, FULL_W)
                two = doc.add_section(WD_SECTION.CONTINUOUS)
                set_columns(two, 2)

    # --- references
    para(doc, [("References", {})], style="Heading 1")
    for key in order:
        p = para(doc, style="references")
        add_runs(p, [(f"[{cite_no[key]}]	", {})] + ieee_reference(bib[key]))

    doc.core_properties.title = "".join(t for t, _ in inline(title.replace("\\\\", " ")))
    doc.core_properties.author = "; ".join("".join(t for t, _ in inline(n)) for n, _ in blocks)
    doc.save(OUT)
    unused = sorted(set(bib) - set(order))
    print(f"wrote {OUT}: {sec_no} sections, {nt} tables, {nf} figure(s), {len(order)} references"
          + (f" (unused bib keys: {unused})" if unused else ""))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
