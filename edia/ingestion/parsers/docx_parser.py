from docx import Document


def parse_docx(file_path: str) -> tuple[str, dict]:
    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # also pull table cell text -- decision docs frequently carry the key facts in a table
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))
    full_text = "\n".join(paragraphs + table_text)
    return full_text, {"paragraph_count": len(paragraphs), "table_count": len(doc.tables)}
